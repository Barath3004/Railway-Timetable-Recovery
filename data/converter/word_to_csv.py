"""
Word to CSV Converter

Purpose
-------
Reads a Word table (.docx) and converts it into CSV.

Usage
-----
python -m data.converter.word_to_csv
"""

from pathlib import Path
from docx import Document
import csv


def convert_docx_to_csv(docx_file: str, csv_file: str):

    print(f"[INFO] Reading {docx_file}")

    document = Document(docx_file)

    if len(document.tables) == 0:
        raise Exception("No table found inside the Word document.")

    table = document.tables[0]

    rows = []

    for row in table.rows:

        values = []

        for cell in row.cells:
            values.append(cell.text.strip())

        rows.append(values)

    Path(csv_file).parent.mkdir(parents=True, exist_ok=True)

    with open(
        csv_file,
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as file:

        writer = csv.writer(file)

        writer.writerows(rows)

    print(f"[SUCCESS] CSV saved to {csv_file}")


if __name__ == "__main__":

    convert_docx_to_csv(
        "data/raw/chennai_division.docx",
        "data/raw/chennai_division.csv"
    )

    convert_docx_to_csv(
        "data/raw/madurai_division.docx",
        "data/raw/madurai_division.csv"
    )

    convert_docx_to_csv(
        "data/raw/thiruvananthapuram_division.docx",
        "data/raw/thiruvananthapuram_division.csv"
    )

    print("\nAll files converted successfully.")